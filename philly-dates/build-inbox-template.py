#!/usr/bin/env python3
"""Generate Philly Dates two-sheet inbox template (Descriptions + Happy Hour Times)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DAY_NAMES = [
    "Sunday",
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
]

OUT = Path(__file__).resolve().parent / "philly-dates-inbox-template.xlsx"

DESC_HEADERS = [
    "Status",
    "App",
    "Name",
    "Address",
    "Neighborhood",
    "Description",
    "Social media links",
    "TikTok link",
    "Website / Menu",
]

TIME_HEADERS = ["Restaurant", "Neighborhood", "Address"]
for day in DAY_NAMES:
    TIME_HEADERS.append(f"{day} start")
    TIME_HEADERS.append(f"{day} end")

EXAMPLE_DESC = [
    "inbox",
    "dates",
    "",
    "",
    "",
    "Optional notes from the video (bar-only HH, vibe, deals you saw)",
    "https://www.instagram.com/reel/…",
    "",
    "",
]

HEADER_FILL = PatternFill("solid", fgColor="E8F5E9")
HEADER_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def autosize(ws, max_width=48):
    for col_idx, column_cells in enumerate(ws.columns, 1):
        letter = get_column_letter(col_idx)
        best = 10
        for cell in column_cells:
            val = cell.value
            if val is None:
                continue
            best = max(best, min(len(str(val)) + 2, max_width))
        ws.column_dimensions[letter].width = best


def main():
    wb = Workbook()
    desc = wb.active
    desc.title = "Descriptions"

    desc.append(DESC_HEADERS)
    for c in range(1, len(DESC_HEADERS) + 1):
        cell = desc.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP

    desc.append(EXAMPLE_DESC)
    desc.freeze_panes = "A2"
    desc.row_dimensions[2].height = 36

    status_dv = DataValidation(
        type="list",
        formula1='"inbox,ready,needs you,done,skip"',
        allow_blank=True,
    )
    status_dv.error = "Pick a status from the list"
    desc.add_data_validation(status_dv)
    status_dv.add("A2:A5000")

    app_dv = DataValidation(
        type="list",
        formula1='"dates,coffee,both"',
        allow_blank=True,
    )
    app_dv.error = "Pick dates, coffee, or both"
    desc.add_data_validation(app_dv)
    app_dv.add("B2:B5000")

    times = wb.create_sheet("Happy Hour Times")
    times.append(TIME_HEADERS)
    for c in range(1, len(TIME_HEADERS) + 1):
        cell = times.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    times.freeze_panes = "A2"
    # One blank row so Google Sheets treats it as a table
    times.append([""] * len(TIME_HEADERS))

    autosize(desc)
    autosize(times, max_width=22)

    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
